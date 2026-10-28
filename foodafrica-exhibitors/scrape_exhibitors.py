"""Scrape every exhibitor from https://foodafrica-expo.com/exhibitors-list into an Excel workbook.

The site renders the complete list into a single HTML table and paginates it in the browser
(DataTables, 20 rows per page), so one download holds every exhibitor on every page.

Usage:
    python3 scrape_exhibitors.py                       # download the live page
    python3 scrape_exhibitors.py --html saved.html     # parse a saved copy instead
"""
import argparse
import collections
import datetime
import re
import sys

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

URL = "https://foodafrica-expo.com/exhibitors-list"
PAGE_SIZE = 20  # DataTables "iDisplayLength" on the site
USER_AGENT = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
# "Hall 4/ F03", "Hall 5 / F03", "Hall2 / -" and the typo "Hal l2/ F71" all occur on the site.
HALL_RE = re.compile(r"^H\s*a\s*l\s*l?\s*(\d+)\s*/\s*(.*)$", re.IGNORECASE)
TEST_ENTRY_RE = re.compile(r"^IFP Test\b", re.IGNORECASE)  # IFP is the organizer

COLUMNS = [
    ("No.", 7),
    ("Company Name", 48),
    ("Country", 20),
    ("Hall", 9),
    ("Stand No.", 13),
    ("Stand (as listed on site)", 22),
    ("Profiles", 55),
    ("Categories", 40),
    ("Logo URL", 40),
    ("Page on Site", 8),
    ("Notes", 42),
]


def clean(text):
    # Collapse whitespace the way a browser displays it.
    return re.sub(r"\s+", " ", text or "").strip()


def badges(node):
    return [clean(b.get_text()) for b in node.select("span.company-info-badge")] if node else []


def popup_value(info, label):
    for h4 in info.find_all("h4"):
        if clean(h4.get_text()).rstrip(":").lower() == label.lower():
            return h4
    return None


def split_stand(stand):
    """Return (hall, stand) parsed from the site's stand text; the original text is kept separately."""
    m = HALL_RE.match(stand)
    if not m:
        return "", stand
    rest = m.group(2).strip()
    return f"Hall {int(m.group(1))}", "" if rest in ("", "-") else rest


def parse(html):
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table", id="exhibitors_list_table")
    if table is None:
        sys.exit("Exhibitor table not found; the page layout may have changed.")
    records, problems = [], []
    for i, tr in enumerate(table.find("tbody").find_all("tr", recursive=False), start=1):
        name_td, country_td, stand_td, profiles_td, categories_td = tr.find_all("td", recursive=False)
        raw_name = name_td.select_one("div.exhibitor-company > span:not(.exhibitor-logo-box)").get_text()
        logo = name_td.select_one("div.exhibitor-company img")
        rec = {
            "no": i,
            "company": clean(raw_name),
            "country": clean(country_td.get_text()),
            "stand_listed": clean(stand_td.get_text()),
            "profiles": badges(profiles_td),
            "categories": badges(categories_td),
            "logo_url": logo.get("src", "").strip() if logo else "",
            "page": (i - 1) // PAGE_SIZE + 1,
            "line_breaks_in_name": "\n" in raw_name.strip(),
        }
        rec["hall"], rec["stand"] = split_stand(rec["stand_listed"])

        # Each value also appears in the row's detail popup; any disagreement or unknown field is reported.
        info = name_td.find("div", class_="company-info")
        popup = {"company": clean(info.find("h3").get_text())}
        for key, label in (("country", "Country"), ("stand_listed", "Stand Number")):
            h4 = popup_value(info, label)
            popup[key] = clean(h4.find_next_sibling("span").get_text()) if h4 else ""
        h4 = popup_value(info, "Profiles")
        popup["profiles"] = badges(h4.find_parent("div")) if h4 else []
        for key, value in popup.items():
            if rec[key] != value:
                problems.append(f"row {i}: {key} differs between list ({rec[key]!r}) and popup ({value!r})")
        unknown = {clean(h.get_text()).rstrip(":") for h in info.find_all("h4")} - {"Country", "Stand Number", "Profiles"}
        if unknown:
            problems.append(f"row {i}: popup has fields this script does not capture: {sorted(unknown)}")
        records.append(rec)
    return records, problems


def notes_for(rec):
    notes = []
    if TEST_ENTRY_RE.search(rec["company"]):
        notes.append("Looks like an organizer test entry; kept because it is on the site's list.")
    if not rec["stand_listed"]:
        notes.append("No stand number listed on the site.")
    if rec["line_breaks_in_name"]:
        notes.append("Name spaced as the site displays it (its source text has line breaks mid-word).")
    return " ".join(notes)


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_workbook(records, path, extracted_on):
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibitors"
    ws.append([name for name, _ in COLUMNS])
    for rec in records:
        ws.append([
            rec["no"],
            rec["company"],
            rec["country"],
            rec["hall"],
            rec["stand"],
            rec["stand_listed"],
            "; ".join(rec["profiles"]),
            "; ".join(rec["categories"]),
            rec["logo_url"],
            rec["page"],
            notes_for(rec),
        ])
    for idx, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    wrap = Alignment(vertical="top", wrap_text=True)
    top = Alignment(vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap if cell.column in (2, 7, 8, 11) else top
            if cell.column == 9 and cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
                cell.alignment = top
    style_header(ws, 1, len(COLUMNS))
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    table = Table(displayName="Exhibitors", ref=f"A1:{get_column_letter(len(COLUMNS))}{len(records) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)

    summary = wb.create_sheet("Summary")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 16
    summary.column_dimensions["C"].width = 70
    summary["A1"] = "Food Africa 2026 – Exhibitors List"
    summary["A1"].font = Font(bold=True, size=14)
    countries = collections.Counter(r["country"] for r in records)
    halls = collections.Counter(r["hall"] or "Not stated" for r in records)
    facts = [
        ("Source", URL),
        ("Extracted on", extracted_on),
        ("Total exhibitors", len(records)),
        ("Countries", len(countries)),
        ("Pages on site (20 per page)", records[-1]["page"] if records else 0),
        ("With profiles / categories", sum(1 for r in records if r["profiles"] or r["categories"])),
        ("With logo", sum(1 for r in records if r["logo_url"])),
    ]
    for offset, (label, value) in enumerate(facts, start=3):
        summary.cell(row=offset, column=1, value=label).font = Font(bold=True)
        summary.cell(row=offset, column=2, value=value).alignment = Alignment(horizontal="left")

    row = len(facts) + 4
    summary.cell(row=row, column=1, value="Notes").font = Font(bold=True)
    notes = [
        "Rows are in the site's own A–Z order; 'Page on Site' is the list page each exhibitor appears on.",
        "'Hall' and 'Stand No.' are split out of the site's stand text, which is kept unchanged in "
        "'Stand (as listed on site)'. A stand shown as '-' on the site is left blank in 'Stand No.'.",
        "Profiles and Categories are the tags each exhibitor selected (some are free text), separated by '; '.",
        "The list page's heading says it covers Halls 2–6.",
    ]
    for offset, note in enumerate(notes, start=row + 1):
        summary.cell(row=offset, column=1, value=note)
        summary.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=3)
        summary.cell(row=offset, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[offset].height = 30

    row = row + len(notes) + 2
    summary.cell(row=row, column=1, value="Hall")
    summary.cell(row=row, column=2, value="Exhibitors")
    style_header(summary, row, 2)
    hall_order = sorted(halls, key=lambda h: (h == "Not stated", h))
    for offset, hall in enumerate(hall_order, start=row + 1):
        summary.cell(row=offset, column=1, value=hall)
        summary.cell(row=offset, column=2, value=halls[hall])

    row = row + len(hall_order) + 2
    summary.cell(row=row, column=1, value="Country")
    summary.cell(row=row, column=2, value="Exhibitors")
    style_header(summary, row, 2)
    for offset, (country, count) in enumerate(sorted(countries.items(), key=lambda kv: (-kv[1], kv[0])), start=row + 1):
        summary.cell(row=offset, column=1, value=country)
        summary.cell(row=offset, column=2, value=count)

    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--html", help="parse this saved copy of the page instead of downloading it")
    parser.add_argument("--out", default="FoodAfrica_2026_Exhibitors.xlsx", help="output .xlsx path")
    args = parser.parse_args()

    if args.html:
        with open(args.html, encoding="utf-8") as f:
            html = f.read()
    else:
        resp = requests.get(URL, headers={"User-Agent": USER_AGENT}, timeout=120)
        resp.raise_for_status()
        resp.encoding = "utf-8"
        html = resp.text

    records, problems = parse(html)
    for problem in problems:
        print("WARNING:", problem, file=sys.stderr)
    write_workbook(records, args.out, datetime.date.today().isoformat())
    print(f"{len(records)} exhibitors written to {args.out}")


if __name__ == "__main__":
    main()
